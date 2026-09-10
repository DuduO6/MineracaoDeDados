"""Create a blinded, stratified Excel sheet for manual sentiment annotation."""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import settings

SEED = 42
DEFAULT_SAMPLE_SIZE = 600
VALIDATION_DIR = settings.ROOT_DIR / "data" / "validation"


def as_bool(series: pd.Series) -> pd.Series:
    return series.astype("string").str.lower().eq("true")


def stratified_sample(data: pd.DataFrame, size: int, seed: int) -> pd.DataFrame:
    """Ensure representation per match/period/prediction, then fill randomly."""
    strata = ["match_id", "comment_period", "sentiment"]
    groups = list(data.groupby(strata, observed=True, dropna=False))
    base_n = max(1, size // max(1, len(groups)))
    pieces = [group.sample(n=min(base_n, len(group)), random_state=seed + index)
              for index, (_, group) in enumerate(groups)]
    sample = pd.concat(pieces) if pieces else data.iloc[0:0]
    if len(sample) < size:
        remaining = data.loc[~data.index.isin(sample.index)]
        sample = pd.concat([sample, remaining.sample(n=min(size - len(sample), len(remaining)), random_state=seed)])
    elif len(sample) > size:
        sample = sample.sample(n=size, random_state=seed)
    return sample.sample(frac=1, random_state=seed).reset_index(drop=True)


def create_workbook(sample: pd.DataFrame, path: Path) -> None:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instruções"
    instructions.append(["VALIDAÇÃO MANUAL DE SENTIMENTO"])
    instructions.append([])
    instructions.append(["Objetivo", "Classificar cada comentário sem consultar a previsão automática."])
    instructions.append(["positive", "Elogio, comemoração, apoio, satisfação ou expectativa positiva."])
    instructions.append(["neutral", "Informação, pergunta, ambiguidade ou ausência de polaridade dominante."])
    instructions.append(["negative", "Crítica, raiva, frustração, hostilidade ou provocação depreciativa."])
    instructions.append(["Casos mistos", "Use a polaridade dominante; se estiver equilibrado, marque neutral."])
    instructions.append(["Sarcasmo", "Classifique o sentido pretendido, não apenas as palavras literais."])
    instructions.append(["Incerto", "Selecione yes quando contexto insuficiente ou interpretação duvidosa."])
    instructions.append(["Privacidade", "Não tente identificar o autor nem inferir o clube do usuário."])
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 100
    instructions["A1"].font = Font(bold=True, size=14)
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet = workbook.create_sheet("Classificação")
    headers = ["sample_id", "data", "partida", "placar", "resultado", "período",
               "comentário", "sentimento_manual", "incerto", "observações"]
    sheet.append(headers)
    for row in sample.itertuples(index=False):
        sheet.append([row.sample_id, row.data, row.partida, row.placar, row.resultado,
                      row.período, row.comentário, "", "no", ""])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(2, sheet.max_row + 1):
        for column in (8, 9, 10):
            sheet.cell(row, column).fill = input_fill
        for column in range(1, 11):
            sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)

    sentiment_validation = DataValidation(type="list", formula1='"positive,neutral,negative"', allow_blank=True)
    uncertain_validation = DataValidation(type="list", formula1='"no,yes"', allow_blank=False)
    sentiment_validation.error = "Selecione positive, neutral ou negative."
    sentiment_validation.errorTitle = "Valor inválido"
    sentiment_validation.prompt = "Escolha a polaridade na lista."
    sentiment_validation.promptTitle = "Sentimento manual"
    sentiment_validation.showInputMessage = True
    sentiment_validation.showErrorMessage = True
    sheet.add_data_validation(sentiment_validation)
    sheet.add_data_validation(uncertain_validation)
    sentiment_validation.add(f"H2:H{sheet.max_row}")
    uncertain_validation.add(f"I2:I{sheet.max_row}")

    colors = {"positive": "C6EFCE", "neutral": "D9EAD3", "negative": "FFC7CE"}
    for label, color in colors.items():
        sheet.conditional_formatting.add(
            f"H2:H{sheet.max_row}", FormulaRule(formula=[f'$H2="{label}"'], fill=PatternFill("solid", fgColor=color))
        )
    sheet.freeze_panes = "G2"
    sheet.auto_filter.ref = f"A1:J{sheet.max_row}"
    widths = {"A": 13, "B": 13, "C": 34, "D": 10, "E": 22, "F": 13,
              "G": 90, "H": 21, "I": 12, "J": 42}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 28
    sheet.sheet_view.showGridLines = False
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--size", type=int, default=DEFAULT_SAMPLE_SIZE)
    parser.add_argument("--seed", type=int, default=SEED)
    args = parser.parse_args()

    nlp = pd.read_csv(settings.PROCESSED_DIR / "comments_nlp.csv", dtype={"comment_id": "string", "match_id": "string"})
    matches = pd.read_csv(settings.RAW_DIR / "matches.csv", dtype={"match_id": "string"})
    eligible = nlp.loc[
        as_bool(nlp["within_comment_window"]) & nlp["comment_period"].isin(["pre_game", "post_game"])
    ].copy()
    sample = stratified_sample(eligible, args.size, args.seed)
    info = matches.set_index("match_id")
    sample["sample_id"] = [f"S{index:04d}" for index in range(1, len(sample) + 1)]
    sample["data"] = sample["match_id"].map(info["match_date"])
    sample["partida"] = sample["match_id"].map(info["home_team"]) + " × " + sample["match_id"].map(info["away_team"])
    sample["placar"] = sample["match_id"].map(info["home_score"]).astype(str) + "–" + sample["match_id"].map(info["away_score"]).astype(str)
    sample["resultado"] = sample["match_id"].map(info["winner"]).replace({"draw": "Empate"})
    sample["período"] = sample["comment_period"].replace({"pre_game": "Pré-jogo", "post_game": "Pós-jogo"})
    sample["comentário"] = sample["text_original"].fillna("")

    population = eligible.groupby(["match_id", "comment_period", "sentiment"], observed=True).size().rename("population_n")
    sampled = sample.groupby(["match_id", "comment_period", "sentiment"], observed=True).size().rename("sample_n")
    key = sample[["sample_id", "comment_id", "match_id", "comment_period", "sentiment", "sentiment_score"]].copy()
    key = key.join(population, on=["match_id", "comment_period", "sentiment"]).join(sampled, on=["match_id", "comment_period", "sentiment"])
    key["sampling_weight"] = key["population_n"] / key["sample_n"]

    workbook_path = VALIDATION_DIR / "sentiment_manual_600.xlsx"
    key_path = VALIDATION_DIR / "sentiment_validation_key.csv"
    create_workbook(sample, workbook_path)
    key.to_csv(key_path, index=False)
    print(f"Planilha cega: {workbook_path}")
    print(f"Chave reservada para avaliação: {key_path}")
    print(f"Comentários selecionados: {len(sample)}")


if __name__ == "__main__":
    main()
